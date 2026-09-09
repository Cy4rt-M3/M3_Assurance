"""Import controls from Excel into MongoDB."""

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pymongo import MongoClient, UpdateOne
from pymongo.collection import Collection

from apps.shared.settings import Settings


def _clean_value(value: Any) -> Any:
    """Convert Excel empty values into MongoDB-safe values."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def read_controls(
    file_path: str,
    sheet_name: str,
) -> list[dict[str, Any]]:
    """Read controls from the Excel worksheet."""
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Controls workbook not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Worksheet '{sheet_name}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise ValueError(f"Worksheet '{sheet_name}' is empty")

        normalized_headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        controls: list[dict[str, Any]] = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            control = {
                header: _clean_value(value)
                for header, value in zip(normalized_headers, row)
                if header
            }

            if control.get("Control ID"):
                controls.append(control)

        return controls
    finally:
        workbook.close()


def create_indexes(collection: Collection[dict[str, Any]]) -> None:
    """Create indexes for the controls collection."""
    collection.create_index(
        [
            ("control_id", 1),
            ("framework", 1),
            ("version", 1),
        ],
        unique=True,
        name="control_framework_version_unique",
    )

    collection.create_index(
        "control_id",
        name="control_id_index",
    )


def import_controls(
    collection: Collection[dict[str, Any]],
    controls: list[dict[str, Any]],
) -> int:
    """Idempotently upsert controls into MongoDB."""
    if not controls:
        return 0

    operations: list[UpdateOne] = []

    for control in controls:
        control_id = control["Control ID"]
        framework = control.get("Framework")
        version = control.get("Version")

        filter_query = {
            "control_id": control_id,
            "framework": framework,
            "version": version,
        }

        document = {
            **control,
            "control_id": control_id,
            "framework": framework,
            "version": version,
        }

        operations.append(
            UpdateOne(
                filter_query,
                {"$set": document},
                upsert=True,
            )
        )

    result = collection.bulk_write(operations, ordered=False)

    return result.upserted_count + result.modified_count


def run_import(settings: Settings | None = None) -> int:
    """Run the complete Excel to MongoDB import."""
    config = settings or Settings()

    controls = read_controls(
        config.controls_file,
        config.controls_sheet,
    )

    client = MongoClient(config.mongodb_url)

    try:
        database = client[config.mongodb_database]
        collection = database[config.mongodb_collection]

        create_indexes(collection)
        changed = import_controls(collection, controls)

        print(
            f"[controls-importer] Read {len(controls)} controls "
            f"from {config.controls_file}"
        )
        print(
            f"[controls-importer] MongoDB: "
            f"{config.mongodb_database}.{config.mongodb_collection}"
        )
        print(
            f"[controls-importer] Upserted/updated {changed} controls"
        )

        return len(controls)
    finally:
        client.close()


if __name__ == "__main__":
    run_import()
